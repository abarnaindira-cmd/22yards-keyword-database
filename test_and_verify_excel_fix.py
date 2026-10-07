import sys
import os
import openpyxl
import pandas as pd

sys.path.insert(0, r"d:\22yards_keyword_database_competitor_ready\22yards_keyword_database")

from app.database import SessionLocal
from app.services.excel_exporter import export_harvested_keywords_to_excel
from export_mysql_to_excel import export_to_excel

def main():
    db = SessionLocal()
    target_asin = "T2YVIVA000017"

    print("=" * 95)
    print(f" STEP 1: EXPORTING HARVESTED KEYWORDS EXCEL REPORT FOR ASIN: {target_asin}")
    print("=" * 95)

    excel_path = export_harvested_keywords_to_excel(db, source_product_asin=target_asin)
    print(f"Exported Excel File: {excel_path}\n")

    print("=" * 95)
    print(" STEP 2: PROGRAMMATICAL EXCEL CELL VERIFICATION (READING VIA OPENPYXL)")
    print("=" * 95)

    wb = openpyxl.load_workbook(excel_path)
    if "Harvested Keywords" not in wb.sheetnames:
        print("FAIL: Sheet 'Harvested Keywords' not found in workbook.")
        db.close()
        return

    sheet = wb["Harvested Keywords"]
    rows = list(sheet.iter_rows(values_only=True))

    header = rows[0]
    print(f"Header Row (Columns A-J):")
    print(f"  A: {header[0]}")
    print(f"  B: {header[1]}")
    print(f"  C: {header[2]}")
    print(f"  D: {header[3]}")
    print(f"  E: {header[4]}")
    print(f"  F: {header[5]}")
    print(f"  G: {header[6]}")
    print(f"  H: {header[7]}")
    print(f"  I: {header[8]}")
    print(f"  J: {header[9]}")
    print("-" * 95)

    data_rows = rows[1:]
    print(f"Data Rows Count: {len(data_rows)}\n")

    non_blank_urls = 0
    all_rows_valid = True

    for r_idx, row in enumerate(data_rows, 1):
        asin = row[0]
        pname = row[1]
        kw = row[2]
        source = row[3]
        rel = row[4]
        u1, u2, u3, u4, u5 = row[5], row[6], row[7], row[8], row[9]

        urls_in_row = [u for u in [u1, u2, u3, u4, u5] if u and str(u).strip().startswith("https://www.amazon.in/dp/")]
        non_blank_urls += len(urls_in_row)

        print(f"Row {r_idx:2d} | ASIN: {asin} | Keyword: '{kw}' | Valid URLs: {len(urls_in_row)}/5")
        print(f"       Col F (URL 1): {u1}")
        print(f"       Col G (URL 2): {u2}")
        print(f"       Col H (URL 3): {u3}")
        print(f"       Col I (URL 4): {u4}")
        print(f"       Col J (URL 5): {u5}")
        print("-" * 95)

        if len(urls_in_row) < 5:
            all_rows_valid = False

    print("\n" + "=" * 95)
    print(" STEP 3: TESTING MAIN EXCEL EXPORTER (marketlens_export.xlsx)")
    print("=" * 95)

    export_to_excel()
    main_export_path = r"d:\22yards_keyword_database_competitor_ready\22yards_keyword_database\marketlens_export.xlsx"

    wb_main = openpyxl.load_workbook(main_export_path)
    sheet_main = wb_main["Harvested Keywords"]
    main_rows = list(sheet_main.iter_rows(values_only=True))
    main_data = [r for r in main_rows[1:] if r[0] == target_asin]

    print(f"\nMain Export File ({main_export_path}) Sheet 'Harvested Keywords':")
    print(f"Rows found for ASIN {target_asin}: {len(main_data)}")
    for m_idx, m_row in enumerate(main_data, 1):
        m_urls = [u for u in m_row[5:10] if u and str(u).strip().startswith("https://www.amazon.in/dp/")]
        print(f"  Row {m_idx:2d} | Keyword: '{m_row[2]}' | URLs Populated: {len(m_urls)}/5")

    print("\n" + "=" * 95)
    print(" FINAL VERIFICATION REPORT FOR EXCEL FIX")
    print("=" * 95)
    print(f"- Target ASIN Tested:        {target_asin}")
    print(f"- Keyword Rows in Excel:     {len(data_rows)}")
    print(f"- Total Non-Blank URLs:      {non_blank_urls} / 50 expected URLs")
    print(f"- Header Column Titles F-J:  {header[5:10]}")
    if all_rows_valid and non_blank_urls == 50:
        print("- VERIFICATION RESULT:       SUCCESS! Columns F-J are 100% populated with stored Amazon URLs.")
    else:
        print("- VERIFICATION RESULT:       FAIL - Some cells are blank.")
    print("=" * 95 + "\n")

    db.close()

if __name__ == "__main__":
    main()
