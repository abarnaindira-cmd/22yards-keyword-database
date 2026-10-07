import openpyxl
from sqlalchemy import text
from app.database import SessionLocal
from app.services.title_generator import process_final_product_data_for_asin
from app.services.excel_exporter import export_final_product_data_to_excel

def test_single_product():
    db = SessionLocal()
    target_sku = "T2YVIVA000017"

    print("==========================================================================")
    print(f"STEP 9: TESTING FINAL PRODUCT DATA FOR SINGLE PRODUCT ({target_sku})")
    print("==========================================================================")

    # 1. Process single product
    res = process_final_product_data_for_asin(db, target_sku)
    print("Processed Result Dict:")
    print(f"  sku_id             : {res.get('sku_id')}")
    print(f"  product_name       : {res.get('product_name')}")
    print(f"  final_product_title: {res.get('final_product_title')}")
    print(f"  keywords_count     : {res.get('keywords_count')}")

    # 2. Inspect database record in MySQL
    db_rec = db.execute(text("SELECT product_name, sku_id, final_product_title, all_keywords FROM final_product_data WHERE sku_id = :sku"), {"sku": target_sku}).fetchone()

    print("\nVerified MySQL Database Record in `final_product_data`:")
    print(f"  product_name       : {db_rec[0]}")
    print(f"  sku_id             : {db_rec[1]}")
    print(f"  final_product_title: {db_rec[2]}")
    print(f"  all_keywords (JSON): {db_rec[3]}")

    assert db_rec[1] == target_sku, f"Expected sku_id {target_sku}, got {db_rec[1]}"
    assert db_rec[0] is not None and db_rec[0] != "", "product_name should not be empty"
    assert db_rec[2] is not None and db_rec[2] != "", "final_product_title should not be empty"

    # 3. Export to Excel and inspect openpyxl cells
    print("\nExporting to Excel file...")
    excel_path = export_final_product_data_to_excel(db, sku_id=target_sku)
    wb = openpyxl.load_workbook(excel_path)

    print(f"Generated Excel Filepath: {excel_path}")
    print(f"Sheet Names             : {wb.sheetnames}")
    assert len(wb.sheetnames) == 1, f"Expected exactly 1 worksheet, got {len(wb.sheetnames)}"
    assert wb.sheetnames[0] == "Final Product Data", f"Expected worksheet name 'Final Product Data', got {wb.sheetnames[0]}"

    sheet = wb["Final Product Data"]
    headers = [sheet.cell(row=1, column=c).value for c in range(1, 5)]
    print(f"Excel Headers (Cols A-D): {headers}")
    assert headers == ['product_name', 'sku_id', 'final_product_title', 'all_keywords'], f"Headers mismatch: {headers}"

    r2_vals = [sheet.cell(row=2, column=c).value for c in range(1, 5)]
    print("\nExcel Row 2 Cell Values:")
    print(f"  Col A (product_name)       : {r2_vals[0]}")
    print(f"  Col B (sku_id)             : {r2_vals[1]}")
    print(f"  Col C (final_product_title): {r2_vals[2]}")
    print(f"  Col D (all_keywords)       : {r2_vals[3][:100]}... (Total len: {len(r2_vals[3])})")

    assert r2_vals[1] == target_sku
    assert r2_vals[0] and r2_vals[2] and r2_vals[3]

    db.close()
    print("\n==========================================================================")
    print("SINGLE PRODUCT TEST PASSED 100% SUCCESSFULLY!")
    print("==========================================================================")

if __name__ == "__main__":
    test_single_product()
