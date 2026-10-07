import os
import time
import urllib.request
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

from sqlalchemy import text
from app.database import SessionLocal
from app.models.product import Product
from app.services.title_generator import process_final_product_data_for_asin

def run_batch_tests():
    db = SessionLocal()

    print("==========================================================================")
    print("STEP 1: PREPARING BATCH 2 GROQ AI TITLES (PRODUCTS 101-200 / IDs 104-203)")
    print("==========================================================================")

    b2_prods = db.query(Product).order_by(Product.id.asc()).offset(100).limit(100).all()
    b2_asins = [p.asin for p in b2_prods if p.asin]

    print(f"Batch 2 Product Count : {len(b2_prods)}")
    print(f"Batch 2 First SKU     : {b2_asins[0]} (ID {b2_prods[0].id})")
    print(f"Batch 2 Last SKU      : {b2_asins[-1]} (ID {b2_prods[-1].id})")

    # Generate Groq AI titles for any Batch 2 products that need title generation
    groq_gen_count = 0
    for idx, asin in enumerate(b2_asins, 1):
        res = process_final_product_data_for_asin(db, asin)
        if res.get("generation_method") == "groq":
            groq_gen_count += 1
        time.sleep(1.0)

    print(f"Batch 2 Groq Titles Status : {groq_gen_count}/100 generated cleanly.")

    print("\n==========================================================================")
    print("STEP 2: TESTING LIVE EXPORT FOR BATCH 1 (PRODUCTS 1-100 / IDs 4-103)")
    print("==========================================================================")

    b1_prods = db.query(Product).order_by(Product.id.asc()).offset(0).limit(100).all()
    b1_asins = [p.asin for p in b1_prods if p.asin]

    # Test GET http://localhost:8001/api/products/export-final-product-data?batch_index=1
    url_b1 = "http://localhost:8001/api/products/export-final-product-data?batch_index=1"
    req_b1 = urllib.request.Request(url_b1)
    path_b1 = os.path.join("scratch", "Batch_1_Export_Test.xlsx")

    with urllib.request.urlopen(req_b1) as resp:
        data_b1 = resp.read()
    with open(path_b1, "wb") as f:
        f.write(data_b1)

    wb_b1 = openpyxl.load_workbook(path_b1)
    sheet_b1 = wb_b1.active
    rows_b1 = sheet_b1.max_row
    cols_b1 = sheet_b1.max_column
    headers_b1 = [sheet_b1.cell(row=1, column=c).value for c in range(1, cols_b1 + 1)]
    skus_b1 = [sheet_b1.cell(row=r, column=2).value for r in range(2, rows_b1 + 1) if sheet_b1.cell(row=r, column=2).value]

    print(f"  - Download URL               : {url_b1}")
    print(f"  - File Path                  : {path_b1}")
    print(f"  - File Size                  : {os.path.getsize(path_b1)} bytes")
    print(f"  - Total Rows                 : {rows_b1} (1 Header + {rows_b1 - 1} Data Rows)")
    print(f"  - Total Columns              : {cols_b1}")
    print(f"  - Headers                    : {headers_b1}")
    print(f"  - First Exported SKU         : {skus_b1[0]} (Expected: {b1_asins[0]})")
    print(f"  - Last Exported SKU          : {skus_b1[-1]} (Expected: {b1_asins[-1]})")

    assert rows_b1 == 101, f"Batch 1 expected 101 rows, got {rows_b1}"
    assert cols_b1 == 4, f"Batch 1 expected 4 columns, got {cols_b1}"
    assert headers_b1 == ["product_name", "sku_id", "final_product_title", "all_keywords"]
    assert skus_b1[0] == b1_asins[0], f"First SKU mismatch: {skus_b1[0]} vs {b1_asins[0]}"
    assert skus_b1[-1] == b1_asins[-1], f"Last SKU mismatch: {skus_b1[-1]} vs {b1_asins[-1]}"
    assert set(skus_b1) == set(b1_asins), "Batch 1 SKU set mismatch!"
    assert len(skus_b1) == 100, f"Batch 1 SKU count mismatch: {len(skus_b1)}"

    print(">>> BATCH 1 TEST PASSED 100% PERFECTLY!")

    print("\n==========================================================================")
    print("STEP 3: TESTING LIVE EXPORT FOR BATCH 2 (PRODUCTS 101-200 / IDs 104-203)")
    print("==========================================================================")

    url_b2 = "http://localhost:8001/api/products/export-final-product-data?batch_index=2"
    req_b2 = urllib.request.Request(url_b2)
    path_b2 = os.path.join("scratch", "Batch_2_Export_Test.xlsx")

    with urllib.request.urlopen(req_b2) as resp:
        data_b2 = resp.read()
    with open(path_b2, "wb") as f:
        f.write(data_b2)

    wb_b2 = openpyxl.load_workbook(path_b2)
    sheet_b2 = wb_b2.active
    rows_b2 = sheet_b2.max_row
    cols_b2 = sheet_b2.max_column
    headers_b2 = [sheet_b2.cell(row=1, column=c).value for c in range(1, cols_b2 + 1)]
    skus_b2 = [sheet_b2.cell(row=r, column=2).value for r in range(2, rows_b2 + 1) if sheet_b2.cell(row=r, column=2).value]

    print(f"  - Download URL               : {url_b2}")
    print(f"  - File Path                  : {path_b2}")
    print(f"  - File Size                  : {os.path.getsize(path_b2)} bytes")
    print(f"  - Total Rows                 : {rows_b2} (1 Header + {rows_b2 - 1} Data Rows)")
    print(f"  - Total Columns              : {cols_b2}")
    print(f"  - Headers                    : {headers_b2}")
    print(f"  - First Exported SKU         : {skus_b2[0]} (Expected: {b2_asins[0]})")
    print(f"  - Last Exported SKU          : {skus_b2[-1]} (Expected: {b2_asins[-1]})")

    assert rows_b2 == 101, f"Batch 2 expected 101 rows, got {rows_b2}"
    assert cols_b2 == 4, f"Batch 2 expected 4 columns, got {cols_b2}"
    assert headers_b2 == ["product_name", "sku_id", "final_product_title", "all_keywords"]
    assert skus_b2[0] == b2_asins[0], f"First SKU mismatch: {skus_b2[0]} vs {b2_asins[0]}"
    assert skus_b2[-1] == b2_asins[-1], f"Last SKU mismatch: {skus_b2[-1]} vs {b2_asins[-1]}"
    assert set(skus_b2) == set(b2_asins), "Batch 2 SKU set mismatch!"
    assert len(skus_b2) == 100, f"Batch 2 SKU count mismatch: {len(skus_b2)}"

    print(">>> BATCH 2 TEST PASSED 100% PERFECTLY!")

    db.close()

    print("\n==========================================================================")
    print("ALL BATCH COMPARISON & EXPORT VERIFICATION TESTS PASSED SUCCESSFULLY!")
    print("==========================================================================")

if __name__ == "__main__":
    run_batch_tests()
