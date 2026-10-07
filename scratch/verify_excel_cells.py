import openpyxl
import io
import os
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

print("==========================================================================")
print("1. VERIFYING SINGLE ASIN EXCEL EXPORT (T2YVIVA000017)")
print("==========================================================================")
resp = client.get("/api/keywords/export-harvested-keywords?asin=T2YVIVA000017&format=excel")
assert resp.status_code == 200, f"Expected 200 OK, got {resp.status_code}"

wb = openpyxl.load_workbook(io.BytesIO(resp.content))
sheet = wb["Harvested Keywords"]

print(f"Sheet Name: {sheet.title}")
headers = [sheet.cell(row=1, column=c).value for c in range(1, 11)]
print(f"Headers (Cols A-J): {headers}")

print("\nCell values for rows 2 to 6:")
for r in range(2, 7):
    asin_val = sheet.cell(row=r, column=1).value
    kw_val = sheet.cell(row=r, column=3).value
    f_val = sheet.cell(row=r, column=6).value
    g_val = sheet.cell(row=r, column=7).value
    h_val = sheet.cell(row=r, column=8).value
    i_val = sheet.cell(row=r, column=9).value
    j_val = sheet.cell(row=r, column=10).value
    print(f"Row {r}: ASIN={asin_val} | Keyword={kw_val}")
    print(f"   Cell F{r} (URL 1): {f_val}")
    print(f"   Cell G{r} (URL 2): {g_val}")
    print(f"   Cell H{r} (URL 3): {h_val}")
    print(f"   Cell I{r} (URL 4): {i_val}")
    print(f"   Cell J{r} (URL 5): {j_val}")
    assert f_val and f_val.startswith("http"), f"Cell F{r} is blank or invalid!"

print("\n==========================================================================")
print("2. VERIFYING ALL PRODUCTS EXCEL EXPORT (Harvested_Keywords_All_Products.xlsx)")
print("==========================================================================")
resp_all = client.get("/api/keywords/export-harvested-keywords?format=excel")
assert resp_all.status_code == 200, f"Expected 200 OK, got {resp_all.status_code}"

wb_all = openpyxl.load_workbook(io.BytesIO(resp_all.content))
sheet_all = wb_all["Harvested Keywords"]

print(f"Sheet Name: {sheet_all.title}")
headers_all = [sheet_all.cell(row=1, column=c).value for c in range(1, 11)]
print(f"Headers (Cols A-J): {headers_all}")

print("\nCell values for rows 2 to 6 in All Products Export:")
for r in range(2, 7):
    asin_val = sheet_all.cell(row=r, column=1).value
    kw_val = sheet_all.cell(row=r, column=3).value
    f_val = sheet_all.cell(row=r, column=6).value
    g_val = sheet_all.cell(row=r, column=7).value
    h_val = sheet_all.cell(row=r, column=8).value
    i_val = sheet_all.cell(row=r, column=9).value
    j_val = sheet_all.cell(row=r, column=10).value
    print(f"Row {r}: ASIN={asin_val} | Keyword={kw_val}")
    print(f"   Cell F{r} (URL 1): {f_val}")
    print(f"   Cell G{r} (URL 2): {g_val}")
    print(f"   Cell H{r} (URL 3): {h_val}")
    print(f"   Cell I{r} (URL 4): {i_val}")
    print(f"   Cell J{r} (URL 5): {j_val}")
    assert f_val and f_val.startswith("http"), f"Cell F{r} is blank or invalid in All Products Export!"

# Check total populated URLs across the entire sheet
total_rows = sheet_all.max_row
populated_url1 = sum(1 for r in range(2, total_rows + 1) if sheet_all.cell(row=r, column=6).value and str(sheet_all.cell(row=r, column=6).value).startswith("http"))
print(f"\nTotal rows in Excel: {total_rows}")
print(f"Total rows with non-empty URL 1 in Column F: {populated_url1}")

print("\n==========================================================================")
print("ALL VERIFICATIONS PASSED PERFECTLY!")
print("==========================================================================")
