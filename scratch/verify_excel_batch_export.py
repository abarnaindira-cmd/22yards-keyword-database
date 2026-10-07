import os
import openpyxl
from app.database import SessionLocal
from app.models.product import Product
from app.services.excel_exporter import export_final_product_data_to_excel

db = SessionLocal()

products = db.query(Product).filter(Product.id.between(4, 103)).all()
target_asins = [p.asin for p in products if p.asin]

test_filepath = export_final_product_data_to_excel(db, asins_list=target_asins)
filename = os.path.basename(test_filepath)

print(f"Saved Excel export to: {test_filepath}")
print(f"File Name: {filename}")
print(f"File Size: {os.path.getsize(test_filepath)} bytes")

# Read generated Excel with openpyxl
wb = openpyxl.load_workbook(test_filepath)
sheet = wb.active

print(f"Sheet Title: {sheet.title}")
print(f"Total Rows : {sheet.max_row}")
print(f"Total Cols : {sheet.max_column}")

headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
print("Headers:", headers)

assert headers == ["product_name", "sku_id", "final_product_title", "all_keywords"], f"Mismatch headers: {headers}"
assert sheet.max_row == 101, f"Expected 101 rows (1 header + 100 products), got {sheet.max_row}"

row_2 = [sheet.cell(row=2, column=c).value for c in range(1, 5)]
print("\nRow 2 Sample Data:")
print("  product_name       :", row_2[0])
print("  sku_id             :", row_2[1])
print("  final_product_title:", row_2[2])
print("  all_keywords       :", row_2[3])

db.close()
print("\n==========================================================================")
print("EXCEL EXPORT FOR UPLOADED BATCH VERIFIED 100% SUCCESSFULLY!")
print("==========================================================================")
