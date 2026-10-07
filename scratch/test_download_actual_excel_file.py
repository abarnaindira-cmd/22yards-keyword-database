import os
import urllib.request
import openpyxl

print("==========================================================================")
print("TESTING ACTUAL DOWNLOADED EXCEL FILE FROM FASTAPI ENDPOINT")
print("==========================================================================")

api_url = "http://localhost:8001/api/products/export-final-product-data"
test_output_path = os.path.join("scratch", "Downloaded_Final_Product_Data_Test.xlsx")

# 1. Download file from live endpoint
print(f"1. Requesting GET: {api_url}")
req = urllib.request.Request(api_url)

with urllib.request.urlopen(req) as response:
    content_type = response.headers.get("Content-Type")
    content_disposition = response.headers.get("Content-Disposition")
    data = response.read()

print(f"   Response Status     : 200 OK")
print(f"   Content-Type        : {content_type}")
print(f"   Content-Disposition : {content_disposition}")
print(f"   Downloaded Size     : {len(data)} bytes")

with open(test_output_path, "wb") as f:
    f.write(data)

# 2. Inspect actual downloaded Excel file
print("\n2. Inspecting Downloaded Excel File with openpyxl...")
wb = openpyxl.load_workbook(test_output_path)
sheet_names = wb.sheetnames
print(f"   Worksheet Names    : {sheet_names}")

sheet = wb.active
total_rows = sheet.max_row
total_cols = sheet.max_column

print(f"   Total Rows in File : {total_rows} (Expected: 101 -> 1 Header + 100 Product Rows)")
print(f"   Total Columns      : {total_cols} (Expected: 4)")

headers = [sheet.cell(row=1, column=c).value for c in range(1, total_cols + 1)]
print(f"   Headers            : {headers}")

# 3. Verify exact assertions
assert len(sheet_names) == 1, f"Expected 1 worksheet, got {len(sheet_names)}"
assert headers == ["product_name", "sku_id", "final_product_title", "all_keywords"], f"Header mismatch: {headers}"
assert total_rows == 101, f"CRITICAL FAILURE: Expected 101 total rows (100 products + 1 header), but found {total_rows} rows!"

# 4. Verify sample row data
row_2 = [sheet.cell(row=2, column=c).value for c in range(1, 5)]
print("\n3. Sample Product Data Row 2:")
print(f"   Column 1 (product_name)       : {row_2[0]}")
print(f"   Column 2 (sku_id)             : {row_2[1]}")
print(f"   Column 3 (final_product_title): {row_2[2]}")
print(f"   Column 4 (all_keywords)       : {row_2[3][:100]}...")

assert row_2[0] and len(str(row_2[0])) > 0
assert row_2[1] and len(str(row_2[1])) > 0
assert row_2[2] and len(str(row_2[2])) > 0
assert row_2[3] and isinstance(row_2[3], str) and "," in row_2[3]

print("\n==========================================================================")
print("SUCCESS: DOWNLOADED EXCEL FILE VERIFIED! EXACTLY 101 ROWS (100 PRODUCTS + 1 HEADER)")
print("==========================================================================")
