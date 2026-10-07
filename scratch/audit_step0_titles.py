"""
STEP 0 - READ-ONLY title audit for the confirmed SKU list (22yards).

What it does (nothing else):
  * Reads products, final_product_data, competitor_products and keywords
    for the given SKUs from MySQL using SELECT statements only, inside a
    READ ONLY transaction (MySQL itself rejects any write).
  * Snapshots the current final_product_data rows to JSON (rollback point).
  * Classifies each existing final_product_title.
  * Reports competitor-title counts, master-facts coverage and
    duplicate name / colour / size information.

What it never does:
  * No INSERT / UPDATE / DELETE / DDL. It does NOT import app.database
    (that module runs CREATE DATABASE IF NOT EXISTS on import).
  * No Groq calls, no scraping, no network access other than MySQL.
  * Does not modify any project file. Output goes to a new folder:
        scratch/audit_step0_<timestamp>/

Run from the project root (22yards_keyword_database):
    python scratch/audit_step0_titles.py
Options:
    --skus-file scratch/skus_100.txt
    --master    22Y_SKU_Master_Updated.xlsx
    --out-dir   scratch/audit_step0_<timestamp>
"""
import argparse
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, date
from difflib import SequenceMatcher
from decimal import Decimal

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MANUAL = "MANUAL_REVIEW_REQUIRED"
EMPTY_VALUES = {"", "NA", "N/A", "NAN", "NONE", "NULL", "-"}

# Claim / marketing / use-case / audience words. In an existing title they are
# flagged only when they do NOT appear in our product name or verified facts.
CLAIM_WORDS = {
    "perfect", "ideal", "suitable", "best", "premium", "quality", "durable", "durability",
    "ultimate", "superior", "advanced", "professional", "pro", "recreational", "beginner",
    "beginners", "comfortable", "comfort", "lightweight", "waterproof", "anti", "fog", "uv",
    "protection", "resistant", "leakproof", "ergonomic", "performance", "training", "practice",
    "competition", "tournament", "match", "pool", "open", "outdoor", "indoor", "gym",
    "kids", "kid", "children", "boys", "girls", "women", "womens", "ladies", "men", "mens",
    "adult", "adults", "unisex", "junior", "senior", "silicone", "breathable", "quick", "dry",
}

AUDIT_CSV_COLUMNS = [
    "sku_id", "in_products_table", "product_id", "product_name_db", "category_db", "product_status",
    "fpd_row_exists", "existing_title", "title_length", "title_class", "title_flags",
    "keywords_count", "fpd_all_keywords_count",
    "competitor_rows", "competitor_distinct_titles", "competitor_placeholder_titles",
    "competitor_keywords_used", "competitor_title_bucket", "max_similarity_to_competitor",
    "master_found", "master_brand", "master_category", "master_colour", "master_size",
    "master_invoice_description", "master_website_name", "website_name_matches_db_name",
    "master_facts_present",
    "name_group_size_in_batch", "name_group_size_in_db", "name_colour_size_group_size_in_batch",
    "variant_differentiator", "ambiguous_after_colour_size",
]


# ---------------------------------------------------------------- helpers
def clean(text):
    if text is None:
        return ""
    s = str(text)
    s = (s.replace("‑", "-").replace("‐", "-").replace("–", "-")
          .replace("—", "-").replace("‘", "'").replace("’", "'")
          .replace("“", '"').replace("”", '"'))
    return " ".join(re.sub(r"[^\x00-\x7F]+", " ", s).split())


def norm(text):
    return " ".join(re.findall(r"[a-z0-9]+", clean(text).casefold()))


def fact(value):
    v = clean(value)
    return "" if v.upper() in EMPTY_VALUES else v


def numbers_in(text):
    return set(re.findall(r"\d+(?:\.\d+)?", clean(text)))


def load_env(path):
    values = {}
    try:
        from dotenv import dotenv_values
        values = {k: v for k, v in dotenv_values(path).items() if v is not None}
    except Exception:
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "=" in line:
                        k, v = line.split("=", 1)
                        values[k.strip()] = v.strip().strip('"').strip("'")
    for k in ("DB_HOST", "DB_PORT", "DB_USER", "DB_PASSWORD", "DB_NAME"):
        if os.getenv(k):
            values.setdefault(k, os.getenv(k))
    return values


def load_skus(path):
    with open(path, encoding="utf-8") as f:
        raw = re.split(r"[,\s]+", f.read())
    seen, out = set(), []
    for s in raw:
        s = s.strip()
        if s and s not in seen:
            seen.add(s)
            out.append(s)
    return out


def load_master(path):
    """sku -> facts dict from the first master row for that SKU."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    cols = {
        "brand": "Brand Name", "category": "Product Category", "colour": "Colour",
        "size": "Size", "invoice_description": "Product Description on Invoice",
        "website_name": "Name On Brand WebSite",
    }
    sku_i = idx["SKU ID for Market Places"]
    out, dup_counts = {}, Counter()
    for r in rows:
        sku = clean(r[sku_i]) if r[sku_i] is not None else ""
        if not sku:
            continue
        dup_counts[sku] += 1
        if sku in out:
            continue
        out[sku] = {k: (fact(r[idx[c]]) if c in idx else "") for k, c in cols.items()}
    wb.close()
    return out, dup_counts


# --------------------------------------------------------- read-only database
class ReadOnlyDB:
    """SELECT-only wrapper around a DB-API connection in a READ ONLY transaction."""

    def __init__(self, conn, placeholder="%s", mysql=True):
        self.conn = conn
        self.ph = placeholder
        cur = conn.cursor()
        if mysql:
            cur.execute("SET SESSION TRANSACTION READ ONLY")
            cur.execute("START TRANSACTION READ ONLY")
        cur.close()

    def select(self, sql, params=()):
        if not sql.lstrip().upper().startswith("SELECT"):
            raise RuntimeError("Audit refused a non-SELECT statement.")
        if self.ph != "%s":
            sql = sql.replace("%s", self.ph)
        cur = self.conn.cursor()
        cur.execute(sql, tuple(params))
        cols = [d[0] for d in cur.description]
        rows = [dict(zip(cols, row)) for row in cur.fetchall()]
        cur.close()
        return rows

    def close(self):
        try:
            self.conn.rollback()
        finally:
            self.conn.close()


def connect_mysql(env):
    import pymysql
    conn = pymysql.connect(
        host=env.get("DB_HOST", "localhost"),
        port=int(env.get("DB_PORT", 3306)),
        user=env.get("DB_USER", "root"),
        password=env.get("DB_PASSWORD", ""),
        database=env.get("DB_NAME", "marketlens"),
        charset="utf8mb4",
        autocommit=False,
    )
    return ReadOnlyDB(conn)


def in_clause(values):
    return "(" + ", ".join(["%s"] * len(values)) + ")"


def fetch_data(db, skus):
    q = in_clause(skus)
    products = db.select(
        f"SELECT id, asin, product_name, category, status FROM products WHERE asin IN {q}", skus)
    fpd = db.select(
        "SELECT product_name, sku_id, final_product_title, all_keywords "
        f"FROM final_product_data WHERE sku_id IN {q}", skus)
    comps = db.select(
        "SELECT source_product_asin, keyword, competitor_asin, competitor_rank, competitor_title "
        f"FROM competitor_products WHERE source_product_asin IN {q} "
        "ORDER BY source_product_asin, competitor_rank", skus)
    kw_counts = db.select(
        "SELECT source_product_asin, COUNT(*) AS n FROM keywords "
        f"WHERE source_product_asin IN {q} GROUP BY source_product_asin", skus)

    names = sorted({norm(p["product_name"]) for p in products if p.get("product_name")})
    # DB-wide count of products sharing each name (what the current code checks).
    all_names = db.select("SELECT product_name FROM products")
    db_name_counts = Counter(norm(r["product_name"]) for r in all_names if r.get("product_name"))
    return products, fpd, comps, kw_counts, {n: db_name_counts[n] for n in names}


# ------------------------------------------------------------ classification
def classify_title(title, product_name):
    """Return the title class used to decide regeneration later."""
    if title is None:
        return "EMPTY"
    t = clean(title)
    if not t:
        return "EMPTY"
    if norm(t) == norm(MANUAL):
        return "MANUAL"
    name = clean(product_name)
    if norm(t) == norm(name):
        return "SAME_AS_NAME"
    if "|" in t:
        if "|" not in name:
            return "LEGACY_FALLBACK"
        # Our own name contains '|': legacy only if the name is followed by another '|' block.
        if norm(t).startswith(norm(name)) and t.count("|") > name.count("|"):
            return "LEGACY_FALLBACK"
    return "EXISTING_TITLE"


def title_flags(title, product_name, facts, competitor_titles, is_variant_group):
    """Pre-checks on an EXISTING_TITLE. Informational only; nothing is regenerated."""
    flags = []
    t = clean(title)
    source = " ".join([clean(product_name)] + [facts.get(k, "") for k in
                      ("brand", "category", "colour", "size", "invoice_description", "website_name")])
    src_tokens = set(norm(source).split())
    t_tokens = norm(t).split()

    unverified = sorted({w for w in t_tokens if w in CLAIM_WORDS and w not in src_tokens})
    if unverified:
        flags.append("unverified_claim:" + "/".join(unverified))
    invented = sorted(numbers_in(t) - numbers_in(source))
    if invented:
        flags.append("invented_number:" + "/".join(invented))
    brand = norm(facts.get("brand", ""))
    if brand and brand not in norm(t):
        flags.append("brand_missing")
    if is_variant_group:
        colour, size = norm(facts.get("colour", "")), norm(facts.get("size", ""))
        if colour and colour not in norm(t):
            flags.append("variant_colour_missing")
        if size and size.replace(" ", "") not in norm(t).replace(" ", ""):
            flags.append("variant_size_missing")
    if "|" in t and "|" not in clean(product_name):
        flags.append("contains_pipe")
    if len(t) > 200:
        flags.append("too_long")
    elif len(t) < 40:
        flags.append("short_title")
    if re.search(r"\b(\w+)\s+\1\b", t, flags=re.I):
        flags.append("repeated_word")
    best = max((SequenceMatcher(None, norm(t), norm(c)).ratio() for c in competitor_titles), default=0.0)
    if best >= 0.85:
        flags.append("near_competitor_copy")
    return flags, round(best, 3)


def comp_bucket(n):
    if n == 0:
        return "0"
    if n < 5:
        return "1-4"
    if n < 10:
        return "5-9"
    if n < 20:
        return "10-19"
    return "20+"


def json_default(o):
    if isinstance(o, (datetime, date)):
        return o.isoformat()
    if isinstance(o, Decimal):
        return float(o)
    if isinstance(o, bytes):
        return o.decode("utf-8", "replace")
    return str(o)


# ------------------------------------------------------------------- audit
def build_audit(skus, master, master_dups, products, fpd, comps, kw_counts, db_name_counts):
    prod_by = {p["asin"]: p for p in products}
    fpd_by = {r["sku_id"]: r for r in fpd}
    kw_by = {r["source_product_asin"]: int(r["n"]) for r in kw_counts}
    comps_by = defaultdict(list)
    for c in comps:
        comps_by[c["source_product_asin"]].append(c)

    def name_of(sku):
        p = prod_by.get(sku)
        return clean(p["product_name"]) if p else master.get(sku, {}).get("website_name", "")

    name_groups = Counter(norm(name_of(s)) for s in skus)
    ncs_key = lambda s: (norm(name_of(s)), norm(master.get(s, {}).get("colour", "")),
                         norm(master.get(s, {}).get("size", "")))
    ncs_groups = Counter(ncs_key(s) for s in skus)

    rows = []
    for sku in skus:
        p = prod_by.get(sku)
        f = master.get(sku, {})
        name = name_of(sku)
        r = fpd_by.get(sku)
        title = r["final_product_title"] if r else None

        comp_rows = comps_by.get(sku, [])
        titles, placeholders = [], 0
        for c in comp_rows:
            t = clean(c.get("competitor_title"))
            if not t:
                continue
            if re.fullmatch(r"Product [A-Z0-9]{10}", t):
                placeholders += 1
                continue
            if t not in titles:
                titles.append(t)
        kw_used = len({c.get("keyword") for c in comp_rows if c.get("keyword")})

        cls = "NO_FPD_ROW" if r is None else classify_title(title, name)
        group_n = name_groups[norm(name)]
        flags, best_sim = ([], "")
        if cls == "EXISTING_TITLE":
            flags, best_sim = title_flags(title, name, f, titles, group_n > 1)
            cls = "EXISTING_TITLE_FLAGGED" if flags else "EXISTING_TITLE_PASSES_PRECHECK"

        try:
            fpd_kw = r["all_keywords"] if r else None
            if isinstance(fpd_kw, (bytes, str)):
                fpd_kw = json.loads(fpd_kw)
            fpd_kw_n = len(fpd_kw) if isinstance(fpd_kw, list) else ""
        except Exception:
            fpd_kw_n = "unparseable"

        differ = []
        if group_n > 1:
            if f.get("colour"):
                differ.append("colour")
            if f.get("size"):
                differ.append("size")
        facts_present = [k for k in ("brand", "category", "colour", "size", "invoice_description", "website_name") if f.get(k)]

        rows.append({
            "sku_id": sku,
            "in_products_table": "yes" if p else "NO",
            "product_id": p["id"] if p else "",
            "product_name_db": clean(p["product_name"]) if p else "",
            "category_db": clean(p["category"]) if p else "",
            "product_status": p["status"] if p else "",
            "fpd_row_exists": "yes" if r else "no",
            "existing_title": clean(title) if title else "",
            "title_length": len(clean(title)) if title else 0,
            "title_class": cls,
            "title_flags": "; ".join(flags),
            "keywords_count": kw_by.get(sku, 0),
            "fpd_all_keywords_count": fpd_kw_n,
            "competitor_rows": len(comp_rows),
            "competitor_distinct_titles": len(titles),
            "competitor_placeholder_titles": placeholders,
            "competitor_keywords_used": kw_used,
            "competitor_title_bucket": comp_bucket(len(titles)),
            "max_similarity_to_competitor": best_sim,
            "master_found": "yes" if sku in master else "NO",
            "master_brand": f.get("brand", ""),
            "master_category": f.get("category", ""),
            "master_colour": f.get("colour", ""),
            "master_size": f.get("size", ""),
            "master_invoice_description": f.get("invoice_description", ""),
            "master_website_name": f.get("website_name", ""),
            "website_name_matches_db_name": ("yes" if norm(f.get("website_name")) == norm(p["product_name"]) else "no") if (p and f) else "",
            "master_facts_present": ",".join(facts_present),
            "name_group_size_in_batch": group_n,
            "name_group_size_in_db": db_name_counts.get(norm(name), 0),
            "name_colour_size_group_size_in_batch": ncs_groups[ncs_key(sku)],
            "variant_differentiator": ",".join(differ) if group_n > 1 else "n/a (unique name)",
            "ambiguous_after_colour_size": "YES" if ncs_groups[ncs_key(sku)] > 1 else "no",
            "_master_dup_rows": master_dups.get(sku, 0),
        })
    return rows


def summarize(rows, skus, master_path, db_name):
    lines = []
    add = lines.append
    n = len(rows)
    add("=" * 78)
    add(f"STEP 0 READ-ONLY TITLE AUDIT  |  {datetime.now():%Y-%m-%d %H:%M:%S}")
    add(f"Database: {db_name}  |  SKUs audited: {n}  |  Master: {os.path.basename(master_path)}")
    add("=" * 78)

    add("\n[1] Presence")
    add(f"  In products table        : {sum(r['in_products_table'] == 'yes' for r in rows)}/{n}")
    add(f"  final_product_data row   : {sum(r['fpd_row_exists'] == 'yes' for r in rows)}/{n}")
    add(f"  Found in master Excel    : {sum(r['master_found'] == 'yes' for r in rows)}/{n}")
    missing = [r["sku_id"] for r in rows if r["in_products_table"] != "yes"]
    if missing:
        add(f"  NOT in products table    : {missing}")

    add("\n[2] Existing title classification")
    for cls, c in Counter(r["title_class"] for r in rows).most_common():
        add(f"  {cls:32s}: {c}")
    flag_counter = Counter()
    for r in rows:
        for fl in filter(None, r["title_flags"].split("; ")):
            flag_counter[fl.split(":")[0]] += 1
    if flag_counter:
        add("  Flags on existing titles (a title can have several):")
        for fl, c in flag_counter.most_common():
            add(f"    - {fl:28s}: {c}")

    add("\n[3] Competitor titles (distinct, placeholders excluded)")
    for b in ("0", "1-4", "5-9", "10-19", "20+"):
        add(f"  {b:6s}: {sum(r['competitor_title_bucket'] == b for r in rows)}")
    ph = sum(r["competitor_placeholder_titles"] for r in rows)
    add(f"  'Product <ASIN>' placeholder titles found: {ph} "
        f"(on {sum(r['competitor_placeholder_titles'] > 0 for r in rows)} SKUs)")
    zero = [r["sku_id"] for r in rows if r["competitor_distinct_titles"] == 0]
    if zero:
        add(f"  SKUs with NO competitor titles: {zero}")

    add("\n[4] Master-facts coverage")
    for k in ("brand", "category", "colour", "size", "invoice_description", "website_name"):
        add(f"  {k:20s}: {sum(bool(r['master_' + k]) for r in rows)}/{n}")
    mism = sum(r["website_name_matches_db_name"] == "no" for r in rows)
    add(f"  Master website name differs from products.product_name: {mism}")
    md = [r["sku_id"] for r in rows if r["_master_dup_rows"] > 1]
    add(f"  SKUs with >1 row in master: {md if md else 'none'}")

    add("\n[5] Duplicate names / colour / size")
    groups = defaultdict(list)
    for r in rows:
        groups[norm(r["product_name_db"] or r["master_website_name"])].append(r)
    dup_groups = {k: v for k, v in groups.items() if len(v) > 1}
    add(f"  Shared-name groups in batch: {len(dup_groups)}  "
        f"(SKUs involved: {sum(len(v) for v in dup_groups.values())})")
    add(f"  SKUs still ambiguous after name+colour+size: "
        f"{sum(r['ambiguous_after_colour_size'] == 'YES' for r in rows)}")
    wider = [r["sku_id"] for r in rows if r["name_group_size_in_db"] > r["name_group_size_in_batch"]]
    add(f"  SKUs whose name is shared with products OUTSIDE this batch: {len(wider)}")
    for _, members in sorted(dup_groups.items(), key=lambda kv: -len(kv[1])):
        name = members[0]["product_name_db"] or members[0]["master_website_name"]
        add(f"\n  \"{name}\"  ({len(members)} SKUs, {members[0]['name_group_size_in_db']} in whole DB)")
        for m in members:
            add(f"     {m['sku_id']:18s} colour={m['master_colour'] or '-':14s} size={m['master_size'] or '-':10s} "
                f"title_class={m['title_class']}")
    add("")
    return "\n".join(lines)


def write_outputs(out_dir, rows, fpd_rows, skus, summary, env_db):
    os.makedirs(out_dir, exist_ok=True)
    csv_path = os.path.join(out_dir, "audit_report.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=AUDIT_CSV_COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    snap_path = os.path.join(out_dir, "final_product_data_snapshot.json")
    snapshot = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "database": env_db,
        "table": "final_product_data",
        "sku_count_requested": len(skus),
        "rows_found": len(fpd_rows),
        "skus_without_row": [s for s in skus if s not in {r["sku_id"] for r in fpd_rows}],
        "rows": fpd_rows,
    }
    with open(snap_path, "w", encoding="utf-8") as f:
        json.dump(snapshot, f, indent=2, ensure_ascii=False, default=json_default)

    sum_path = os.path.join(out_dir, "audit_summary.txt")
    with open(sum_path, "w", encoding="utf-8") as f:
        f.write(summary)
    return csv_path, snap_path, sum_path


def main(argv=None, db=None):
    ap = argparse.ArgumentParser(description="Step 0 read-only title audit")
    ap.add_argument("--skus-file", default=os.path.join(PROJECT_ROOT, "scratch", "skus_100.txt"))
    ap.add_argument("--master", default=os.path.join(PROJECT_ROOT, "22Y_SKU_Master_Updated.xlsx"))
    ap.add_argument("--out-dir", default=None)
    args = ap.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    skus = load_skus(args.skus_file)
    if not skus:
        print("No SKUs found in", args.skus_file)
        return 1
    master, master_dups = load_master(args.master)

    env = load_env(os.path.join(PROJECT_ROOT, ".env"))
    db_name = env.get("DB_NAME", "marketlens")
    own_db = db is None
    if own_db:
        db = connect_mysql(env)
    try:
        products, fpd, comps, kw_counts, db_name_counts = fetch_data(db, skus)
    finally:
        if own_db:
            db.close()  # rollback + close; nothing was written

    rows = build_audit(skus, master, master_dups, products, fpd, comps, kw_counts, db_name_counts)
    summary = summarize(rows, skus, args.master, db_name)
    out_dir = args.out_dir or os.path.join(PROJECT_ROOT, "scratch", f"audit_step0_{datetime.now():%Y%m%d_%H%M%S}")
    csv_path, snap_path, sum_path = write_outputs(out_dir, rows, fpd, skus, summary, db_name)

    print(summary)
    print("Files written (outside the database):")
    print("  ", csv_path)
    print("  ", snap_path)
    print("  ", sum_path)
    print("Database writes: NONE (read-only transaction, SELECT statements only). Groq calls: NONE.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
