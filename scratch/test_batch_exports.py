import os
import sys
import urllib.request
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

from app.database import SessionLocal
from app.models.product import Product

def test_batch_1_and_batch_2():
    db = SessionLocal()

    # 1. Load Expected SKUs for Batch 1 (IDs 4 to 103)
    b1_prods = db.query(Product).order_by(Product.id.asc()).offset(0).limit(100).all()
    b1_expected_skus = [p.asin for p in b1_prods if p.asin]

    # 2. Load Expected SKUs for Batch 2 (IDs 104 to 203)
    b2_prods = db.query(Product).order_by(Product.id.asc()).offset(100).limit(100).all()
    b2_expected_skus = [p.asin for p in b2_prods if p.asin]

    db.close()

    print("==========================================================================")
    print("RUNNING LIVE EXPORT VERIFICATION TEST FOR BATCH 1 & BATCH 2")
    print("==========================================================================")

    # --------------------------------------------------------------------------
    # BATCH 1 TEST
    # --------------------------------------------------------------------------
    print("\n--- TESTING BATCH 1 (100 PRODUCTS: IDs 4 to 103) ---")
    url_b1 = "http://localhost:8001/api/products/export-final-product-data?batch_index=1"
    path_b1 = os.path.join("scratch", "Live_Export_Batch_1.xlsx")

    with urllib.request.urlopen(url_b1) as resp:
        data_b1 = resp.read()
    with open(path_b1, "wb") as f:
        f.write(data_b1)

    wb_b1 = openpyxl.load_workbook(path_b1)
    sheet_b1 = wb_b1.active
    rows_b1 = sheet_b1.max_row
    cols_b1 = sheet_b1.max_column
    headers_b1 = [sheet_b1.cell(row=1, column=c).value for c in range(1, cols_b1 + 1)]
    skus_b1 = [sheet_b1.cell(row=r, column=2).value for r in range(2, rows_b1 + 1) if sheet_b1.cell(row=r, column=2).value]

    print(f"URL: {url_b1}")
    print(f"File Size: {len(data_b1)} bytes")
    print(f"Total Rows (inc. header): {rows_b1}")
    print(f"Total Columns: {cols_b1}")
    print(f"Headers: {headers_b1}")
    print(f"First Exported SKU: {skus_b1[0]}")
    print(f"Last Exported SKU : {skus_b1[-1]}")
    print(f"Expected First SKU: {b1_expected_skus[0]}")
    print(f"Expected Last SKU : {b1_expected_skus[-1]}")

    # Assertions for Batch 1
    assert rows_b1 == 101, f"Batch 1 expected 101 rows, got {rows_b1}"
    assert cols_b1 == 4, f"Batch 1 expected 4 columns, got {cols_b1}"
    assert headers_b1 == ["product_name", "sku_id", "final_product_title", "all_keywords"], f"Batch 1 headers mismatch: {headers_b1}"
    assert len(skus_b1) == 100, f"Batch 1 expected 100 SKUs, got {len(skus_b1)}"
    assert skus_b1[0] == b1_expected_skus[0], f"Batch 1 first SKU mismatch: {skus_b1[0]} vs {b1_expected_skus[0]}"
    assert skus_b1[-1] == b1_expected_skus[-1], f"Batch 1 last SKU mismatch: {skus_b1[-1]} vs {b1_expected_skus[-1]}"

    b1_missing = set(b1_expected_skus) - set(skus_b1)
    b1_extra = set(skus_b1) - set(b1_expected_skus)
    print(f"Batch 1 Missing SKUs: {len(b1_missing)}, Extra SKUs: {len(b1_extra)}")
    assert len(b1_missing) == 0 and len(b1_extra) == 0, "Batch 1 SKU set mismatch!"

    print(">>> BATCH 1 LIVE EXPORT VERIFICATION PASSED PERFECTLY!")

    # --------------------------------------------------------------------------
    # BATCH 2 TEST
    # --------------------------------------------------------------------------
    print("\n--- TESTING BATCH 2 (100 PRODUCTS: IDs 104 to 203) ---")
    url_b2 = "http://localhost:8001/api/products/export-final-product-data?batch_index=2"
    path_b2 = os.path.join("scratch", "Live_Export_Batch_2.xlsx")

    with urllib.request.urlopen(url_b2) as resp:
        data_b2 = resp.read()
    with open(path_b2, "wb") as f:
        f.write(data_b2)

    wb_b2 = openpyxl.load_workbook(path_b2)
    sheet_b2 = wb_b2.active
    rows_b2 = sheet_b2.max_row
    cols_b2 = sheet_b2.max_column
    headers_b2 = [sheet_b2.cell(row=1, column=c).value for c in range(1, cols_b2 + 1)]
    skus_b2 = [sheet_b2.cell(row=r, column=2).value for r in range(2, rows_b2 + 1) if sheet_b2.cell(row=r, column=2).value]

    print(f"URL: {url_b2}")
    print(f"File Size: {len(data_b2)} bytes")
    print(f"Total Rows (inc. header): {rows_b2}")
    print(f"Total Columns: {cols_b2}")
    print(f"Headers: {headers_b2}")
    print(f"First Exported SKU: {skus_b2[0]}")
    print(f"Last Exported SKU : {skus_b2[-1]}")
    print(f"Expected First SKU: {b2_expected_skus[0]}")
    print(f"Expected Last SKU : {b2_expected_skus[-1]}")

    # Assertions for Batch 2
    assert rows_b2 == 101, f"Batch 2 expected 101 rows, got {rows_b2}"
    assert cols_b2 == 4, f"Batch 2 expected 4 columns, got {cols_b2}"
    assert headers_b2 == ["product_name", "sku_id", "final_product_title", "all_keywords"], f"Batch 2 headers mismatch: {headers_b2}"
    assert len(skus_b2) == 100, f"Batch 2 expected 100 SKUs, got {len(skus_b2)}"
    assert skus_b2[0] == b2_expected_skus[0], f"Batch 2 first SKU mismatch: {skus_b2[0]} vs {b2_expected_skus[0]}"
    assert skus_b2[-1] == b2_expected_skus[-1], f"Batch 2 last SKU mismatch: {skus_b2[-1]} vs {b2_expected_skus[-1]}"

    b2_missing = set(b2_expected_skus) - set(skus_b2)
    b2_extra = set(skus_b2) - set(b2_expected_skus)
    print(f"Batch 2 Missing SKUs: {len(b2_missing)}, Extra SKUs: {len(b2_extra)}")
    assert len(b2_missing) == 0 and len(b2_extra) == 0, "Batch 2 SKU set mismatch!"

    print(">>> BATCH 2 LIVE EXPORT VERIFICATION PASSED PERFECTLY!")

    # --------------------------------------------------------------------------
    # TESTING BATCH 2 VIA EXPLICIT ASINS LIST (SIMULATING UI SELECTION/UPLOAD)
    # --------------------------------------------------------------------------
    print("\n--- TESTING BATCH 2 VIA ASINS QUERY PARAMETER ---")
    asins_param = ",".join(b2_expected_skus)
    url_b2_asins = f"http://localhost:8001/api/products/export-final-product-data?asins={asins_param}"
    path_b2_asins = os.path.join("scratch", "Live_Export_Batch_2_ASINS.xlsx")

    req = urllib.request.Request(url_b2_asins)
    with urllib.request.urlopen(req) as resp:
        data_b2_asins = resp.read()
    with open(path_b2_asins, "wb") as f:
        f.write(data_b2_asins)

    wb_b2_a = openpyxl.load_workbook(path_b2_asins)
    sheet_b2_a = wb_b2_a.active
    rows_b2_a = sheet_b2_a.max_row
    cols_b2_a = sheet_b2_a.max_column
    headers_b2_a = [sheet_b2_a.cell(row=1, column=c).value for c in range(1, cols_b2_a + 1)]
    skus_b2_a = [sheet_b2_a.cell(row=r, column=2).value for r in range(2, rows_b2_a + 1) if sheet_b2_a.cell(row=r, column=2).value]

    assert rows_b2_a == 101, f"Expected 101 rows, got {rows_b2_a}"
    assert cols_b2_a == 4, f"Expected 4 columns, got {cols_b2_a}"
    assert headers_b2_a == ["product_name", "sku_id", "final_product_title", "all_keywords"]
    assert skus_b2_a[0] == b2_expected_skus[0]
    assert skus_b2_a[-1] == b2_expected_skus[-1]
    assert set(skus_b2_a) == set(b2_expected_skus)

    print(">>> BATCH 2 VIA EXPLICIT ASINS PARAMETER PASSED PERFECTLY!")

    print("\n==========================================================================")
    print("ALL VERIFICATION TESTS PASSED 100% PERFECTLY WITH ZERO MISMATCHES!")
    print("==========================================================================")

if __name__ == "__main__":
    test_batch_1_and_batch_2()
