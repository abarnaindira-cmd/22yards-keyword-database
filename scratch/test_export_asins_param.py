import urllib.request
import openpyxl

print("Testing export with specific asins parameter...")
url = "http://localhost:8001/api/products/export-final-product-data?asins=T2YVIVA000017,T2YEVERLAST000005"
req = urllib.request.Request(url)

with urllib.request.urlopen(req) as resp:
    data = resp.read()

filepath = "scratch/test_param_export.xlsx"
with open(filepath, "wb") as f:
    f.write(data)

wb = openpyxl.load_workbook(filepath)
sheet = wb.active
print(f"Total Rows for 2 ASINs: {sheet.max_row} (Expected: 3 -> 1 Header + 2 Products)")
headers = [sheet.cell(row=1, column=c).value for c in range(1, 5)]
print(f"Headers: {headers}")

assert sheet.max_row == 3
assert headers == ["product_name", "sku_id", "final_product_title", "all_keywords"]

print("Param filtering test PASSED successfully!")
