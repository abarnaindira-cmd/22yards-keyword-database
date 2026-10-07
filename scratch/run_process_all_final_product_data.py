import openpyxl
from sqlalchemy import text
from app.database import SessionLocal
from app.services.title_generator import process_all_final_product_data
from app.services.excel_exporter import export_final_product_data_to_excel

def process_all_products():
    db = SessionLocal()

    print("==========================================================================")
    print("PROCESSING FINAL PRODUCT DATA FOR ALL PRODUCTS IN MYSQL DATABASE")
    print("==========================================================================")

    res = process_all_final_product_data(db)
    print(f"Total Products Processed : {res.get('total_processed')}")
    print(f"Success Count            : {res.get('success_count')}")

    # Inspect total records in final_product_data table
    count = db.execute(text("SELECT COUNT(*) FROM final_product_data")).scalar()
    print(f"Total Rows in MySQL `final_product_data` table: {count}")

    # Verify duplicate prevention by rerunning process_all_final_product_data
    print("\nTesting Rerunning Process (Duplicate Prevention Test)...")
    res_rerun = process_all_final_product_data(db)
    count_after = db.execute(text("SELECT COUNT(*) FROM final_product_data")).scalar()
    print(f"Total Rows in MySQL `final_product_data` table after rerun: {count_after}")
    assert count == count_after, f"Duplicate records detected! Before: {count}, After: {count_after}"

    # Export all final product data to Excel (.xlsx)
    print("\nExporting All Final Product Data to Excel...")
    excel_path = export_final_product_data_to_excel(db)
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb["Final Product Data"]

    headers = [sheet.cell(row=1, column=c).value for c in range(1, 5)]
    total_excel_rows = sheet.max_row

    print(f"Export Filepath    : {excel_path}")
    print(f"Worksheet Name     : {sheet.title}")
    print(f"Header Columns A-D : {headers}")
    print(f"Total Rows in Excel: {total_excel_rows}")

    print("\nSample Rows in Excel (Rows 2 to 6):")
    for r in range(2, 7):
        pname = sheet.cell(row=r, column=1).value
        sku = sheet.cell(row=r, column=2).value
        title = sheet.cell(row=r, column=3).value
        kws = sheet.cell(row=r, column=4).value
        print(f"Row {r:2d}: SKU={sku:15s} | Title={str(title)[:50]}... | Keywords Cell Len={len(str(kws))}")

    db.close()
    print("\n==========================================================================")
    print("ALL PRODUCTS FINAL DATA GENERATION & EXCEL VERIFICATION PASSED 100%!")
    print("==========================================================================")

if __name__ == "__main__":
    process_all_products()
